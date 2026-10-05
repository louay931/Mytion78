"""Crée Modele_Demande_EGYDE.docx : formulaire de demande EGYDE avec champs de fusion (source : Fiches_clients.xlsx)."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

BLEU, VALEUR, FOND, GRIS = '1F4D78', '0070C0', 'DEEAF6', 'A6A6A6'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height, s.orientation = Cm(21), Cm(29.7), WD_ORIENT.PORTRAIT
s.left_margin = s.right_margin = Cm(1.5); s.top_margin = s.bottom_margin = Cm(1.3)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
st.paragraph_format.space_after = Pt(0)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    nxt = [e for e in tcPr if e.tag in (qn('w:noWrap'), qn('w:tcMar'), qn('w:textDirection'), qn('w:tcFitText'), qn('w:vAlign'), qn('w:hideMark'))]
    nxt[0].addprevious(sh) if nxt else tcPr.append(sh)  # ordre imposé par le schéma

def borders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), GRIS); b.append(x)
    nxt = [e for e in tblPr if e.tag in (qn('w:shd'), qn('w:tblLayout'), qn('w:tblCellMar'), qn('w:tblLook'))]
    nxt[0].addprevious(b) if nxt else tblPr.append(b)

def champ(p, nom):
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), f' MERGEFIELD {nom} ')
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    for tag, val in (('w:b', None), ('w:color', VALEUR)):
        e = OxmlElement(tag)
        if val: e.set(qn('w:val'), val)
        rPr.append(e)
    t = OxmlElement('w:t'); t.text = f'«{nom}»'; r.append(rPr); r.append(t); f.append(r)
    p._p.append(f)

def ecrire(cell, texte=None, nom=None, gras=False, couleur=None, taille=None, italique=False):
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if texte:
        r = p.add_run(texte); r.bold = gras; r.italic = italique
        if couleur: r.font.color.rgb = RGBColor.from_string(couleur)
        if taille: r.font.size = Pt(taille)
    if nom: champ(p, nom)

def titre_section(texte, note=None):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(texte); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = RGBColor.from_string(BLEU)
    if note:
        r = p.add_run('  ' + note); r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', BLEU)): bt.set(qn(f'w:{k}'), v)
    bd.append(bt)
    sp = pPr.find(qn('w:spacing')); sp.addprevious(bd) if sp is not None else pPr.append(bd)

def tableau(lignes, largeurs):
    t = doc.add_table(rows=len(lignes), cols=len(largeurs)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False; borders(t)
    for i, ligne in enumerate(lignes):
        for j, (case, w) in enumerate(zip(ligne, largeurs)):
            c = t.cell(i, j); c.width = Cm(w)
            if case is None: continue
            if isinstance(case, tuple):   # (libellé)
                ecrire(c, case[0], gras=True); shade(c, 'F2F2F2')
            else:                         # nom de champ
                ecrire(c, nom=case)
    return t

# Titre
p = doc.add_paragraph(); r = p.add_run('DEMANDE EGYDE'); r.bold = True; r.font.size = Pt(20)
r.font.color.rgb = RGBColor.from_string(BLEU)
p = doc.add_paragraph(); r = p.add_run('Intégration / modification d’un site client en télésurveillance')
r.italic = True; r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

titre_section('SITE CLIENT')
tableau([[('Nom du site client',), 'Client', ('Date de la demande',), 'E_DateDemande'],
         [('Adresse',), 'Adresse', ('Date de M.E.S. souhaitée',), 'E_DateMES'],
         [('Code postal',), 'CodePostal', None, None],
         [('Ville',), 'Ville', None, None]], [3.6, 6.2, 4.0, 4.2])
p = doc.add_paragraph(); p.alignment = 2
r = p.add_run('M.E.S. : 3 jours ouvrés minimum si type existant'); r.italic = True; r.font.size = Pt(8.5)
r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

titre_section('CRÉATION', '(à remplir seulement pour la demande d’intégration d’un nouveau site)')
tableau([[('Nom du site',), 'E_NomSite', ('Source principale',), 'E_SourcePrincipale'],
         [('Type de site',), 'E_TypeSite', ('Source attente',), 'E_SourceAttente'],
         [('N° de téléphone',), 'E_TelSite', ('Source secours',), 'E_SourceSecours'],
         [('N° identification',), 'E_NumIdentification', ('Commercial responsable',), 'E_Commercial'],
         [('N° esclave',), 'E_NumEsclave', None, None],
         [('N° astreinte',), 'E_Astreinte', None, None]], [3.6, 6.2, 4.0, 4.2])

titre_section('CONTACTS CLIENT', '(à remplir à la création ou pour l’ajout d’un nouvel utilisateur)')
lib = [('Nom', 'Nom'), ('Prénom', 'Prenom'), ('Adresse mail', 'Mail'), ('N° de téléphone', 'Tel'),
       ('Service', 'Service'), ('Identifiant', 'Identifiant'), ('Mot de passe', 'MotDePasse'),
       ('Type d’accès client', 'TypeAcces'), ('Ajout alarme', 'AjoutAlarme')]
t = tableau([[None, None, None, None]] + [[(l,)] + [f'E_C{n}_{k}' for n in (1, 2, 3)] for l, k in lib],
            [3.6, 4.8, 4.8, 4.8])
for j, txt in enumerate(['', 'Contact 1', 'Contact 2', 'Contact 3']):
    c = t.cell(0, j); shade(c, FOND)
    if txt: ecrire(c, txt, gras=True, couleur=BLEU); c.paragraphs[0].alignment = 1

z = doc.settings.element.find(qn('w:zoom'))  # le modèle python-docx omet w:percent (requis)
if z is not None: z.set(qn('w:percent'), '100')
doc.core_properties.title = 'Demande EGYDE'
doc.save('Modele_Demande_EGYDE.docx'); print('ok')
