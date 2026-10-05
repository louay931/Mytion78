"""Ajoute la demande EGYDE à Fiches_clients.xlsx : section de saisie dans « Saisie » + onglet « Demande EGYDE »
(copie du formulaire EGYDE dont les cases se remplissent depuis « Saisie »).
Usage : python ajout_egyde_excel.py Fiches_clients.xlsx DEMANDE_EGYDE.xlsx sortie.xlsx"""
import copy, sys, openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
src, egyde, out = sys.argv[1], sys.argv[2], sys.argv[3]
wb = openpyxl.load_workbook(src)
ws, fc, lm = wb['Saisie'], wb['Fiches clients'], wb['Lisez-moi']

def style(dst, model):
    for a in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'):
        setattr(dst, a, copy.copy(getattr(model, a)))

TITRE, LIB, VAL, AIDE = ws['A21'], ws['A22'], ws['B22'], ws['C22']
# (libellé, nom du champ Word, aide, type)  type: t=texte, d=date, on=liste oui/non
champs = [('Date de la demande', 'E_DateDemande', 'Format jj/mm/aaaa', 'd'),
          ('Date de M.E.S. souhaitée', 'E_DateMES', '3 jours ouvrés minimum si type existant', 'd'),
          ('Nom du site (création)', 'E_NomSite', 'Ex : CH DE PITHIVIERS / PITHIVIERS (45)', 't'),
          ('Type de site', 'E_TypeSite', 'Ex : LGG', 't'),
          ('N° de téléphone du site', 'E_TelSite', '', 't'),
          ('N° identification', 'E_NumIdentification', '', 't'),
          ('N° esclave', 'E_NumEsclave', '', 't'),
          ('N° astreinte', 'E_Astreinte', '', 't'),
          ('Source principale', 'E_SourcePrincipale', 'Ex : 3000L', 't'),
          ('Source attente', 'E_SourceAttente', 'Ex : 2XCV18', 't'),
          ('Source secours', 'E_SourceSecours', 'Ex : 3B50', 't'),
          ('Commercial responsable', 'E_Commercial', '', 't')]
contact = [('Nom', 'Nom', '', 't'), ('Prénom', 'Prenom', '', 't'), ('Adresse mail', 'Mail', '', 't'),
           ('N° de téléphone', 'Tel', '', 't'), ('Service', 'Service', '', 't'),
           ('Identifiant', 'Identifiant', '', 't'), ('Mot de passe', 'MotDePasse', '', 't'),
           ('Type d’accès client', 'TypeAcces', 'Ex : etendu', 't'), ('Ajout alarme', 'AjoutAlarme', 'oui / non', 'on')]
blocs = [('DEMANDE EGYDE – SITE  (onglet Demande EGYDE ; client, adresse, CP et ville = INFOS CLIENT)', champs)]
for n in (1, 2, 3):
    blocs.append((f'DEMANDE EGYDE – CONTACT CLIENT {n}  (laisser vide si pas de contact {n})',
                  [(l, f'E_C{n}_{k}', a, t) for l, k, a, t in contact]))

assert ws.max_row <= 40, 'la section EGYDE semble déjà présente'
dv_on = DataValidation(type='list', formula1='"oui,non"', allow_blank=True)
ws.add_data_validation(dv_on)
row, lignes_saisie = 41, {}
for titre, lignes in blocs:
    ws.cell(row, 1, titre); style(ws.cell(row, 1), TITRE)
    for c in (2, 3): style(ws.cell(row, c), ws.cell(21, c))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    for lib, nom, aide, typ in lignes:
        a, b, c = ws.cell(row, 1, lib), ws.cell(row, 2), ws.cell(row, 3, aide or None)
        style(a, LIB); style(b, VAL); style(c, AIDE)
        b.number_format = 'DD/MM/YYYY' if typ == 'd' else '@'
        if typ == 'on': dv_on.add(b.coordinate)
        lignes_saisie[nom] = row
        row += 1
    row += 1

# Lisez-moi : point 6 inséré après le point 5
for r in range(1, lm.max_row + 1):
    if str(lm.cell(r, 1).value or '').startswith('5.'):
        lm.insert_rows(r + 1)
        for im in lm._images:  # les captures ne suivent pas insert_rows
            if im.anchor._from.row >= r:
                im.anchor._from.row += 1; im.anchor.to.row += 1
        lm.cell(r + 1, 1, '6. Demande EGYDE : remplis la section « DEMANDE EGYDE » en bas de Saisie, l’onglet « Demande EGYDE » se remplit tout seul '
                          '(nom du site client, adresse, CP et ville repris des INFOS CLIENT). Imprime ou envoie cet onglet.')
        style(lm.cell(r + 1, 1), lm.cell(r, 1))
        break
# Onglet « Demande EGYDE » : copie du formulaire, valeurs remplacées par des liens vers Saisie
src_ws = openpyxl.load_workbook(egyde).active
eg = wb.create_sheet('Demande EGYDE', index=2)
liens = {'B2': 'B4', 'B3': 'B5', 'B5': 'B6', 'B6': 'B7'}  # client, adresse, CP, ville = INFOS CLIENT
cases = {'E_DateDemande': 'G2', 'E_DateMES': 'G3', 'E_NomSite': 'B8', 'E_TypeSite': 'B9', 'E_TelSite': 'B10',
         'E_NumIdentification': 'B11', 'E_NumEsclave': 'B12', 'E_Astreinte': 'B13', 'E_SourcePrincipale': 'G9',
         'E_SourceAttente': 'G10', 'E_SourceSecours': 'G11', 'E_Commercial': 'G12'}
for n, debut in ((1, 16), (2, 26), (3, 36)):
    for i, (_, k, _, _) in enumerate(contact):
        cases[f'E_C{n}_{k}'] = f'B{debut + i}'
liens.update({case: f'B{lignes_saisie[nom]}' for nom, case in cases.items()})
for r in src_ws.iter_rows():
    for c in r:
        d = eg.cell(c.row, c.column)
        style(d, c)
        if c.coordinate in liens:
            ref = f'Saisie!${liens[c.coordinate][0]}${liens[c.coordinate][1:]}'
            d.value = f'=IF({ref}="","",{ref})'
        elif c.coordinate not in ('E3',):  # E3 = espace parasite
            d.value = c.value
for k in ('G2', 'G3'):
    eg[k].number_format = 'DD/MM/YYYY'
for m in src_ws.merged_cells.ranges: eg.merge_cells(str(m))
for k, v in src_ws.column_dimensions.items(): eg.column_dimensions[k].width = v.width
for k, v in src_ws.row_dimensions.items():
    if v.height: eg.row_dimensions[k].height = v.height
for im in src_ws._images: eg.add_image(im)
eg.page_setup.orientation = src_ws.page_setup.orientation
eg.page_setup.paperSize = src_ws.page_setup.paperSize
eg.page_setup.fitToWidth = 1; eg.sheet_properties.pageSetUpPr.fitToPage = True; eg.page_setup.fitToHeight = 0
eg.page_margins = copy.copy(src_ws.page_margins)
eg.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
wb.calculation.fullCalcOnLoad = True
wb.save(out)
print('ok')
